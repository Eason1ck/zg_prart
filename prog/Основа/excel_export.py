# -*- coding: utf-8 -*-
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from data_loader import REGION_NAMES, POINT_KIND_AP, POINT_KIND_AT, POINT_KIND_SP

HEADER_FILL = PatternFill("solid", start_color="305496", end_color="305496")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
DAY_HEADER_FILL = {
    1: PatternFill("solid", start_color="C6E0B4", end_color="C6E0B4"),
    2: PatternFill("solid", start_color="FFE699", end_color="FFE699"),
    3: PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE"),
}
SPECIAL_FILL = PatternFill("solid", start_color="D9D9D9", end_color="D9D9D9")
KIND_FONT = {
    POINT_KIND_AP: Font(name="Arial", bold=True, size=9, color="C00000"),
    POINT_KIND_AT: Font(name="Arial", bold=True, size=9, color="C00000"),
    POINT_KIND_SP: Font(name="Arial", size=9, color="000000"),
}
DEFAULT_FONT = Font(name="Arial", size=9)
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
SPECIAL_LABELS = {
    "price_monitor": "Мониторинг цен",
    "callcenter_monitor": "Мониторинг колл-центра",
    "empty_no_region": "Нет доступного региона",
}


def sanitize_sheet_name(name):
    name = re.sub(r'[\\/?*\[\]:]', "", str(name))
    return name[:31]


def write_employee_sheet(wb, plan, year, month):
    emp = plan.employee
    sheet_name = sanitize_sheet_name(emp.co_name)
    base_name = sheet_name
    suffix = 1
    while sheet_name in wb.sheetnames:
        suffix += 1
        sheet_name = sanitize_sheet_name(f"{base_name} ({suffix})")

    ws = wb.create_sheet(sheet_name)
    max_stops = max((len(d.stops) for d in plan.days), default=1)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4 + max_stops)
    title_cell = ws.cell(row=1, column=1,
                         value=f"{emp.co_name} — {emp.fio} — план на {month:02d}.{year}")
    title_cell.font = Font(name="Arial", bold=True, size=12)

    header_row = 3
    for i, h in enumerate(["Дата", "День недели", "Тип региона", "Кол-во точек"], start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN_BORDER

    for i in range(max_stops):
        c = ws.cell(row=header_row, column=5 + i, value=f"Код {i + 1}")
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN_BORDER

    for i in range(max_stops):
        c = ws.cell(row=header_row, column=5 + max_stops + i, value=f"Точка {i + 1}")
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = THIN_BORDER

    weekday_names = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]

    row = header_row + 1
    for day in plan.days:
        ws.cell(row=row, column=1, value=day.date.strftime("%d.%m.%Y")).border = THIN_BORDER
        ws.cell(row=row, column=2, value=weekday_names[day.date.weekday()]).border = THIN_BORDER

        if day.is_special:
            label = SPECIAL_LABELS.get(day.is_special, day.is_special)
            ws.cell(row=row, column=3, value=label).border = THIN_BORDER
            ws.cell(row=row, column=4, value=0).border = THIN_BORDER
            for col in range(1, 5):
                ws.cell(row=row, column=col).fill = SPECIAL_FILL
        else:
            rt_cell = ws.cell(row=row, column=3, value=REGION_NAMES.get(day.region_type, ""))
            rt_cell.border = THIN_BORDER
            fill = DAY_HEADER_FILL.get(day.region_type)
            if fill:
                rt_cell.fill = fill
            ws.cell(row=row, column=4, value=len(day.stops)).border = THIN_BORDER
            for i, p in enumerate(day.stops):
                # сначала коды
                cc = ws.cell(row=row, column=5 + i, value=p.short_code)
                cc.font = DEFAULT_FONT
                cc.alignment = Alignment(horizontal="center")
                cc.border = THIN_BORDER
            for i, p in enumerate(day.stops):
                # потом названия точек
                c = ws.cell(row=row, column=5 + max_stops + i,
                            value=f"[{p.kind}] {p.name} ({p.region_name})")
                c.font = KIND_FONT.get(p.kind, DEFAULT_FONT)
                c.border = THIN_BORDER

        for col in range(1, 3):
            ws.cell(row=row, column=col).font = DEFAULT_FONT
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")
        row += 1

    row += 2
    warn_title = ws.cell(row=row, column=1, value="Предупреждения и невыполненные требования:")
    warn_title.font = Font(name="Arial", bold=True, size=10, color="C00000")
    row += 1
    if plan.warnings:
        for w in plan.warnings:
            ws.cell(row=row, column=1, value=f"• {w}").font = Font(name="Arial", size=9, color="C00000")
            row += 1
    else:
        ws.cell(row=row, column=1,
                value="Нет замечаний — план выполнен полностью.").font = Font(name="Arial", size=9, color="006100")

    ws.column_dimensions[get_column_letter(1)].width = 12
    ws.column_dimensions[get_column_letter(2)].width = 8
    ws.column_dimensions[get_column_letter(3)].width = 16
    ws.column_dimensions[get_column_letter(4)].width = 12
    for i in range(max_stops):
        ws.column_dimensions[get_column_letter(5 + i)].width = 8
    for i in range(max_stops):
        ws.column_dimensions[get_column_letter(5 + max_stops + i)].width = 32

    ws.freeze_panes = ws.cell(row=header_row + 1, column=5)


def export_plans_to_excel(plans, year, month, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    for plan in plans:
        write_employee_sheet(wb, plan, year, month)
    wb.save(out_path)
    return out_path